import os
import json
import openpyxl
from openpyxl.styles import Font, Alignment, Border, Side

def sync_responses_to_excel(input_json_path=None):
    folder = r"G:\Exium\2026\3Q'26\Survey (Gyne Doctor)"
    master_excel = os.path.join(folder, "Exium_Gyne_Doctor_Survey_Master_2026.xlsx")

    if not input_json_path:
        input_json_path = os.path.join(folder, "survey_responses.json")

    if not os.path.exists(input_json_path):
        print(f"No responses JSON file found at: {input_json_path}")
        return

    with open(input_json_path, "r", encoding="utf-8") as f:
        responses = json.load(f)

    if not responses:
        print("Empty response list.")
        return

    print(f"Loaded {len(responses)} responses from JSON.")

    wb = openpyxl.load_workbook(master_excel)
    ws_resp = wb["Survey Responses"]

    # Collect existing IDs to avoid duplicates (column 22 is Survey Record ID)
    existing_ids = set()
    for r in range(2, ws_resp.max_row + 1):
        v = ws_resp.cell(r, 22).value
        if v:
            existing_ids.add(str(v).strip())

    font_data = Font(name="Calibri", size=10, bold=False, color="FF0F172A")
    c_border = "FFE2E8F0"
    thin_border = Border(
        left=Side(style="thin", color=c_border),
        right=Side(style="thin", color=c_border),
        top=Side(style="thin", color=c_border),
        bottom=Side(style="thin", color=c_border)
    )
    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    start_row = ws_resp.max_row + 1 if ws_resp.max_row >= 2 else 2
    added_count = 0

    for item in responses:
        item_id = str(item.get("id") or "").strip()
        if item_id and item_id in existing_ids:
            continue

        row_vals = [
            item.get("formatted_time") or item.get("timestamp") or "",
            item.get("zone") or "",
            item.get("zonal_head") or "",
            item.get("region") or "",
            item.get("regional_head") or "",
            item.get("sap_territory_code") or "",
            item.get("territory") or "",
            item.get("sap_mio_code") or "",
            item.get("mio_name") or "",
            item.get("doctor_name") or "",
            item.get("doctor_rpl_id") or "",
            item.get("q1_code") or "",
            item.get("q1_answer_en") or "",
            item.get("q2_code") or "",
            item.get("q2_answer_en") or "",
            item.get("q3_code") or "",
            item.get("q3_answer_en") or "",
            item.get("q4_code") or "",
            item.get("q4_answer_en") or "",
            item.get("q5_code") or "",
            item.get("q5_answer_en") or "",
            item_id
        ]

        for col_idx, val in enumerate(row_vals, start=1):
            c = ws_resp.cell(start_row, col_idx, value=val)
            c.font = font_data
            c.border = thin_border
            if col_idx in (1, 6, 8, 11, 12, 14, 16, 18, 20, 22):
                c.alignment = align_center
            else:
                c.alignment = align_left

        start_row += 1
        added_count += 1
        if item_id:
            existing_ids.add(item_id)

    wb.save(master_excel)
    print(f"Successfully synced {added_count} new survey responses to Master Excel.")

if __name__ == '__main__':
    sync_responses_to_excel()
