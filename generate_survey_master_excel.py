import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
import json

def build_master_excel():
    folder = r"G:\Exium\2026\3Q'26\Survey (Gyne Doctor)"
    output_path = os.path.join(folder, "Exium_Gyne_Doctor_Survey_Master_2026.xlsx")
    ff_path = os.path.join(folder, "FF list.xlsx")
    q_config_path = os.path.join(folder, "questions_config.json")

    with open(q_config_path, "r", encoding="utf-8") as f:
        q_cfg = json.load(f)

    wb_ff = openpyxl.load_workbook(ff_path, data_only=True)
    ws_ff = wb_ff["FF, 202609"] if "FF, 202609" in wb_ff.sheetnames else wb_ff.active

    wb = openpyxl.Workbook()
    wb.remove(wb.active) # Remove default sheet

    # Colors & Styles
    c_navy = "FF0F172A"
    c_blue = "FF0284C7"
    c_light_blue = "FFE0F2FE"
    c_green = "FF10B981"
    c_light_green = "FFD1FAE5"
    c_border = "FFE2E8F0"

    fill_navy = PatternFill(start_color=c_navy, end_color=c_navy, fill_type="solid")
    fill_blue = PatternFill(start_color=c_blue, end_color=c_blue, fill_type="solid")
    fill_light_blue = PatternFill(start_color=c_light_blue, end_color=c_light_blue, fill_type="solid")
    fill_light_green = PatternFill(start_color=c_light_green, end_color=c_light_green, fill_type="solid")

    font_title = Font(name="Calibri", size=14, bold=True, color="FFFFFFFF")
    font_subtitle = Font(name="Calibri", size=10, italic=True, color="FFBAE6FD")
    font_section = Font(name="Calibri", size=11, bold=True, color="FF0F172A")
    font_header = Font(name="Calibri", size=10, bold=True, color="FFFFFFFF")
    font_data = Font(name="Calibri", size=10, bold=False, color="FF0F172A")

    thin_border = Border(
        left=Side(style="thin", color=c_border),
        right=Side(style="thin", color=c_border),
        top=Side(style="thin", color=c_border),
        bottom=Side(style="thin", color=c_border)
    )

    align_center = Alignment(horizontal="center", vertical="center")
    align_left = Alignment(horizontal="left", vertical="center")

    # =========================================================================
    # TAB 1: EXECUTIVE SUMMARY
    # =========================================================================
    ws_sum = wb.create_sheet(title="Executive Summary")
    ws_sum.views.sheetView[0].showGridLines = True

    # Title Banner
    ws_sum.merge_cells("A1:I1")
    cell_t1 = ws_sum["A1"]
    cell_t1.value = "EXIUM MUPS — GERD IN PREGNANCY CLINICAL SURVEY REPORT"
    cell_t1.font = font_title
    cell_t1.fill = fill_navy
    cell_t1.alignment = align_center
    ws_sum.row_dimensions[1].height = 32

    ws_sum.merge_cells("A2:I2")
    cell_t2 = ws_sum["A2"]
    cell_t2.value = "Obstetrics & Gynaecology Doctor Opinion Analysis"
    cell_t2.font = font_subtitle
    cell_t2.fill = fill_navy
    cell_t2.alignment = align_center
    ws_sum.row_dimensions[2].height = 18

    total_terrs_count = ws_ff.max_row - 1

    # KPI Overview Cards in Row 4-5
    kpis = [
        ("B4", "C4", "B5", "C5", "TOTAL TERRITORIES", f"{total_terrs_count:,}", fill_light_blue, c_blue),
        ("D4", "E4", "D5", "E5", "TOTAL DOCTORS SURVEYED", "=COUNTA('Survey Responses'!A2:A5000)", fill_light_green, c_green),
        ("G4", "H4", "G5", "H5", "TERRITORY COVERAGE %", f"=(COUNTIF('Territory Coverage Matrix'!J2:J{total_terrs_count+10}, \"Completed\")/{total_terrs_count})*100", fill_light_blue, c_blue),
    ]

    for tl1, br1, tl2, br2, label, formula_val, bg_fill, text_col in kpis:
        ws_sum.merge_cells(f"{tl1}:{br1}")
        ws_sum.merge_cells(f"{tl2}:{br2}")
        c_lbl = ws_sum[tl1]
        c_lbl.value = label
        c_lbl.font = Font(name="Calibri", size=9, bold=True, color="FF475569")
        c_lbl.alignment = align_center
        c_lbl.fill = bg_fill

        c_val = ws_sum[tl2]
        c_val.value = formula_val
        c_val.font = Font(name="Calibri", size=16, bold=True, color="FF0F172A")
        c_val.alignment = align_center
        c_val.fill = bg_fill

    ws_sum.row_dimensions[4].height = 18
    ws_sum.row_dimensions[5].height = 28

    # Dynamic Questions Breakdown (Q1 to Q5)
    curr_row = 7
    for q_idx, q in enumerate(q_cfg["questions"]):
        q_num = q["number"]
        q_title = q["title_en"]
        code_col_letter = get_column_letter(12 + (q_idx * 2)) # Q1: L, Q2: N, Q3: P, Q4: R, Q5: T

        ws_sum.cell(curr_row, 2, value=f"{q_num}. {q_title}").font = font_section

        sub_headers = ["Option Code", "Clinical Response / Option", "Response Count", "% Share"]
        for idx, h in enumerate(sub_headers, start=2):
            cell = ws_sum.cell(curr_row + 2, idx, value=h)
            cell.font = font_header
            cell.fill = fill_blue
            cell.alignment = align_center if idx in (2, 4, 5) else align_left
        ws_sum.row_dimensions[curr_row + 2].height = 22

        opt_start = curr_row + 3
        for opt_idx, opt in enumerate(q["options"]):
            r_idx = opt_start + opt_idx
            r_code = opt["code"]
            r_text = opt["text_en"]

            ws_sum.cell(r_idx, 2, value=r_code).alignment = align_center
            ws_sum.cell(r_idx, 3, value=r_text).alignment = align_left
            ws_sum.cell(r_idx, 4, value=f"=COUNTIF('Survey Responses'!{code_col_letter}:{code_col_letter}, \"{r_code}\")").alignment = align_center
            ws_sum.cell(r_idx, 5, value=f"=IF(COUNTA('Survey Responses'!A2:A5000)>0, D{r_idx}/COUNTA('Survey Responses'!A2:A5000), 0)").alignment = align_center
            ws_sum.cell(r_idx, 5).number_format = "0.0%"

            for c in range(2, 6):
                ws_sum.cell(r_idx, c).font = font_data
                ws_sum.cell(r_idx, c).border = thin_border
            ws_sum.row_dimensions[r_idx].height = 20

        curr_row = opt_start + len(q["options"]) + 2

    ws_sum.column_dimensions["A"].width = 4
    ws_sum.column_dimensions["B"].width = 14
    ws_sum.column_dimensions["C"].width = 52
    ws_sum.column_dimensions["D"].width = 18
    ws_sum.column_dimensions["E"].width = 14
    ws_sum.column_dimensions["F"].width = 14
    ws_sum.column_dimensions["G"].width = 14
    ws_sum.column_dimensions["H"].width = 14
    ws_sum.column_dimensions["I"].width = 4

    # =========================================================================
    # TAB 2: TERRITORY COVERAGE MATRIX
    # =========================================================================
    ws_mat = wb.create_sheet(title="Territory Coverage Matrix")
    ws_mat.views.sheetView[0].showGridLines = True
    ws_mat.freeze_panes = "G2"

    mat_headers = [
        "SL", "Zone", "Zonal Head", "Region", "Regional Head", 
        "SAP Territory Code", "Territory Name", "SAP MIO Code", "MIO / Sr. MIO Name", 
        "Survey Status", "Doctors Surveyed Count"
    ]

    for c_idx, h in enumerate(mat_headers, start=1):
        cell = ws_mat.cell(1, c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_navy
        cell.alignment = align_center
        ws_mat.row_dimensions[1].height = 28

    row_mat = 2
    for r in range(2, ws_ff.max_row + 1):
        terr_code = ws_ff.cell(r, 10).value
        terr_name = ws_ff.cell(r, 8).value
        mio_code = ws_ff.cell(r, 2).value
        mio_name = ws_ff.cell(r, 3).value
        rm_name = ws_ff.cell(r, 13).value
        reg_name = ws_ff.cell(r, 17).value
        zm_name = ws_ff.cell(r, 22).value
        zone_name = ws_ff.cell(r, 26).value

        if not terr_name and not mio_name:
            continue

        ws_mat.cell(row_mat, 1, value=row_mat - 1).alignment = align_center
        ws_mat.cell(row_mat, 2, value=zone_name).alignment = align_left
        ws_mat.cell(row_mat, 3, value=zm_name).alignment = align_left
        ws_mat.cell(row_mat, 4, value=reg_name).alignment = align_left
        ws_mat.cell(row_mat, 5, value=rm_name).alignment = align_left
        ws_mat.cell(row_mat, 6, value=terr_code).alignment = align_center
        ws_mat.cell(row_mat, 7, value=terr_name).alignment = align_left
        ws_mat.cell(row_mat, 8, value=mio_code).alignment = align_center
        ws_mat.cell(row_mat, 9, value=mio_name).alignment = align_left
        
        ws_mat.cell(row_mat, 10, value=f'=IF(K{row_mat}>0, "Completed", "Pending")').alignment = align_center
        ws_mat.cell(row_mat, 11, value=f'=COUNTIF(\'Survey Responses\'!F:F, F{row_mat})').alignment = align_center

        for c in range(1, 12):
            ws_mat.cell(row_mat, c).font = font_data
            ws_mat.cell(row_mat, c).border = thin_border

        row_mat += 1

    mat_widths = {1: 6, 2: 18, 3: 20, 4: 20, 5: 22, 6: 15, 7: 20, 8: 12, 9: 26, 10: 15, 11: 18}
    for col_idx, w in mat_widths.items():
        ws_mat.column_dimensions[get_column_letter(col_idx)].width = w

    # =========================================================================
    # TAB 3: SURVEY RESPONSES (Clean English Raw Response Data - 22 Columns)
    # =========================================================================
    ws_resp = wb.create_sheet(title="Survey Responses")
    ws_resp.views.sheetView[0].showGridLines = True
    ws_resp.freeze_panes = "A2"

    resp_headers = [
        "Timestamp", "Zone", "Zonal Head", "Region", "Regional Head", 
        "SAP Territory Code", "Territory Name", "SAP MIO Code", "MIO / Sr. MIO Name", 
        "Doctor Full Name", "Doctor RPL ID", 
        "Q1 Code", "Q1 Answer (Trimester)", 
        "Q2 Code", "Q2 Answer (GERD Symptom)", 
        "Q3 Code", "Q3 Answer (Lifestyle Resolution)", 
        "Q4 Code", "Q4 Answer (First Choice Medicine)", 
        "Q5 Code", "Q5 Answer (Preferred PPI Molecule)", 
        "Survey Record ID"
    ]

    for c_idx, h in enumerate(resp_headers, start=1):
        cell = ws_resp.cell(1, c_idx, value=h)
        cell.font = font_header
        cell.fill = fill_blue
        cell.alignment = align_center
        ws_resp.row_dimensions[1].height = 32

    resp_widths = {
        1: 20, 2: 18, 3: 20, 4: 20, 5: 22, 6: 15, 7: 20, 8: 12, 9: 26,
        10: 25, 11: 14, 
        12: 10, 13: 26, 
        14: 10, 15: 38, 
        16: 10, 17: 38, 
        18: 10, 19: 34, 
        20: 10, 21: 34, 
        22: 24
    }
    for col_idx, w in resp_widths.items():
        ws_resp.column_dimensions[get_column_letter(col_idx)].width = w

    print(f"Saving Master Excel to: {output_path}")
    wb.save(output_path)
    print("Master Excel successfully built with all 5 questions!")

if __name__ == '__main__':
    build_master_excel()
